"""Column layouts copied from the final (OneDrive) Excel files, so the output can be pasted straight in.

Rebuild after changing columns or tabs in a final file:
    python tools/update_templates.py "path/to/final files folder"
"""
import json
from pathlib import Path

FILE = Path(__file__).parent / "data" / "templates.json"


def load():
    if FILE.exists():
        return json.loads(FILE.read_text(encoding="utf-8"))
    return {}


def headers_of(xlsx_path):
    """{sheet name: [headers]} of a final file."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    out = {}
    for ws in wb:
        first = next(ws.iter_rows(max_row=1, values_only=True), ())
        headers = [str(h).strip() if h is not None else "" for h in first]
        while headers and not headers[-1]:
            headers.pop()
        if headers:
            out[ws.title] = headers
    return out


def varieties_of(xlsx_path):
    """Every variety ever typed in a final file: used to correct words read from pictures."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    out = set()
    for ws in wb:
        rows = ws.iter_rows(values_only=True)
        headers = [str(h).strip().upper() if h is not None else "" for h in next(rows, ())]
        if "VARIETY" not in headers:
            continue
        i = headers.index("VARIETY")
        for r in rows:
            if r and i < len(r) and isinstance(r[i], str) and r[i].strip() and len(r[i]) < 40:
                out.add(r[i].strip())
    return sorted(out)


def build(folder, supplier_list):
    folder = Path(folder)
    data = load()
    found, missing = [], []
    for s in supplier_list:
        f = folder / s.final_file if s.final_file else None
        if f and f.exists():
            data[s.name] = {"file": s.final_file, "sheets": headers_of(f), "varieties": varieties_of(f)}
            found.append(s.name)
        else:
            missing.append(s.name)
    FILE.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    return found, missing
